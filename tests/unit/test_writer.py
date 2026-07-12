# Unit tests for core/writer.py.
#
# These build a throwaway .xlsx in a temporary directory (pytest's `tmp_path`
# fixture), run the async writer functions against it, then re-open the file to
# assert the right cells changed and everything else was preserved.
#
# `asyncio_mode = "auto"` (set in pyproject.toml) lets us write `async def test_*`
# functions without decorating each one — pytest-asyncio runs them on an event loop.

from pathlib import Path

import openpyxl
import pytest

from portfolio.core.exceptions import DuplicateTickerError, PositionNotFoundError
from portfolio.core.models import Position
from portfolio.core.writer import (
    add_position,
    remove_position,
    set_avg_price,
    set_quantity,
)


def _make_workbook(path: Path) -> None:
    """Create a minimal portfolio workbook that matches the reader's layout.

    First sheet: header + two positions. Second sheet: a FixedIncome row we use
    to prove the writer never touches other sheets.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet"
    ws.append(["Ticker", "Quantity", "Exchange", "Category", "Average Price"])
    ws.append(["ITSA4.SA", 100, "B3", "Stock", 7.84])
    ws.append(["WEGE3.SA", 200, "B3", "Stock", 33.05])

    fi = wb.create_sheet("FixedIncome")
    fi.append(["Name", "Amount"])
    fi.append(["CDB Itaú", 604000])
    wb.save(path)


@pytest.fixture()
def book(tmp_path: Path) -> Path:
    """Return the path to a fresh workbook for each test."""
    path = tmp_path / "portfolio.xlsx"
    _make_workbook(path)
    return path


def _rows(path: Path) -> list[tuple]:
    """Read back the first sheet's data rows (skipping the header)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    rows = [tuple(r) for r in ws.iter_rows(min_row=2, values_only=True)]
    wb.close()
    return rows


async def test_set_quantity_updates_only_that_cell(book: Path) -> None:
    await set_quantity(book, "ITSA4.SA", 9105)
    rows = _rows(book)
    assert rows[0] == ("ITSA4.SA", 9105, "B3", "Stock", 7.84)  # qty changed, avg price intact
    assert rows[1] == ("WEGE3.SA", 200, "B3", "Stock", 33.05)  # other row untouched


async def test_set_quantity_matches_despite_whitespace(book: Path) -> None:
    # The reader strips whitespace; the writer must match the same way.
    await set_quantity(book, "  ITSA4.SA  ", 500)
    assert _rows(book)[0][1] == 500


async def test_set_quantity_missing_ticker_raises(book: Path) -> None:
    with pytest.raises(PositionNotFoundError):
        await set_quantity(book, "NOPE.SA", 10)


async def test_set_quantity_rejects_non_positive(book: Path) -> None:
    with pytest.raises(ValueError):
        await set_quantity(book, "ITSA4.SA", 0)


async def test_set_avg_price_updates_only_that_cell(book: Path) -> None:
    await set_avg_price(book, "ITSA4.SA", 6.50)
    rows = _rows(book)
    assert rows[0] == ("ITSA4.SA", 100, "B3", "Stock", 6.50)   # avg price changed, qty intact
    assert rows[1] == ("WEGE3.SA", 200, "B3", "Stock", 33.05)  # other row untouched


async def test_set_avg_price_missing_ticker_raises(book: Path) -> None:
    with pytest.raises(PositionNotFoundError):
        await set_avg_price(book, "NOPE.SA", 10)


async def test_set_avg_price_rejects_non_positive(book: Path) -> None:
    with pytest.raises(ValueError):
        await set_avg_price(book, "ITSA4.SA", 0)


async def test_add_position_appends_row(book: Path) -> None:
    await add_position(
        book,
        Position(ticker="PETR4.SA", quantity=50, exchange="B3", category="Stock", avg_price_native=32.5),
    )
    rows = _rows(book)
    assert len(rows) == 3
    assert rows[2] == ("PETR4.SA", 50, "B3", "Stock", 32.5)


async def test_add_position_without_avg_price_leaves_blank(book: Path) -> None:
    await add_position(book, Position(ticker="MGLU3.SA", quantity=10))
    # The avg-price cell (index 4) should be empty, matching how the reader
    # treats a missing average.
    assert _rows(book)[2][4] is None


async def test_add_duplicate_ticker_raises(book: Path) -> None:
    with pytest.raises(DuplicateTickerError):
        await add_position(book, Position(ticker="ITSA4.SA", quantity=1))


async def test_remove_position(book: Path) -> None:
    await remove_position(book, "ITSA4.SA")
    rows = _rows(book)
    assert len(rows) == 1
    assert rows[0][0] == "WEGE3.SA"  # remaining row shifted up, no gap


async def test_remove_missing_ticker_raises(book: Path) -> None:
    with pytest.raises(PositionNotFoundError):
        await remove_position(book, "NOPE.SA")


async def test_other_sheet_is_preserved(book: Path) -> None:
    # Every edit must leave the FixedIncome sheet completely intact.
    await set_quantity(book, "ITSA4.SA", 1)
    await add_position(book, Position(ticker="BBAS3.SA", quantity=5))
    await remove_position(book, "WEGE3.SA")

    wb = openpyxl.load_workbook(book, data_only=True)
    assert "FixedIncome" in wb.sheetnames
    fi_rows = [tuple(r) for r in wb["FixedIncome"].iter_rows(min_row=2, values_only=True)]
    wb.close()
    assert fi_rows == [("CDB Itaú", 604000)]
