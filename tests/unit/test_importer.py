# Unit tests for core/importer.py — the B3 "posição" import logic.
#
# We synthesise a small B3-format export and a portfolio in a temp directory, then
# assert the parser aggregates correctly and the plan/apply steps do the right
# thing. Building the fixtures ourselves keeps these tests independent of any real
# broker file.

from pathlib import Path

import openpyxl
import pytest

from portfolio.core.importer import (
    _parse_b3_export,
    apply_plan,
    build_plan,
)


def _make_b3_export(path: Path) -> None:
    """Create a miniature B3 export with the multi-sheet shape of the real file.

    Only the columns the parser cares about need to be realistic: a "Produto"
    column at position 0 (whose first token is the ticker) and a "Quantidade"
    column somewhere in the header.
    """
    wb = openpyxl.Workbook()

    acoes = wb.active
    acoes.title = "Acoes"
    acoes.append(["Produto", "Código de Negociação", "Tipo", "Quantidade", "Preço"])
    acoes.append(["ITSA4 - ITAUSA S.A.", "ITSA4", "PN", 9105, 14.17])
    acoes.append(["FLRY3 - FLEURY S.A.", "FLRY3", "ON", 300, 16.42])
    acoes.append(["KLBN11 - KLABIN S.A.", "KLBN11", "UNIT", 300, 17.54])
    acoes.append(["", None, None, None, None])          # blank spacer row
    acoes.append(["", "", "", "", "Total"])             # total row (Produto empty)

    emp = wb.create_sheet("Empréstimos")
    emp.append(["Produto", "Natureza", "Quantidade", "Taxa"])
    emp.append(["FLRY3 - FLEURY S.A.", "Doador", 1760, 0.05])
    emp.append(["BBAS3 - BCO BRASIL S.A.", "Doador", 200, 0.02])   # two BBAS3 rows...
    emp.append(["BBAS3 - BCO BRASIL S.A.", "Doador", 1390, 0.01])  # ...must sum to 1590
    emp.append(["KLBN11 - KLABIN S.A.", "Doador", 2415, 2.69])
    emp.append(["RECR11 - REC RECEBIVEIS", "Doador", 168, 0.14])   # also a Fundo below

    etf = wb.create_sheet("ETF")
    etf.append(["Produto", "Código", "Quantidade"])
    etf.append(["B5P211 - IT NOW IMA-B5", "B5P211", 147])

    fundo = wb.create_sheet("Fundo de Investimento")
    fundo.append(["Produto", "Código", "Quantidade"])
    fundo.append(["RECR11 - REC RECEBIVEIS", "RECR11", 105])   # + 168 lent = 273, Fund
    fundo.append(["RECR12 - REC RECEBIVEIS", "RECR12", 47])    # right; not in portfolio

    # This sheet must be ignored entirely (no market ticker).
    rf = wb.create_sheet("Renda Fixa")
    rf.append(["Produto", "Emissor", "Valor"])
    rf.append(["CDB - ITAU", "ITAU", 12345])

    wb.save(path)


def _make_portfolio(path: Path) -> None:
    """Create a portfolio sheet in the layout the reader expects."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet"
    ws.append(["Ticker", "Quantity", "Exchange", "Category", "Average Price"])
    ws.append(["ITSA4.SA", 9105, "B3", "Stock", 7.84])    # unchanged
    ws.append(["FLRY3.SA", 1760, "B3", "Stock", 15.09])   # -> 2060
    ws.append(["KLBN11.SA", 1515, "B3", "Stock", 18.01])  # -> 2715
    ws.append(["BBAS3.SA", 1000, "B3", "Stock", 23.87])   # -> 1590
    ws.append(["B5P211.SA", 100, "B3", "ETF", 106.83])    # -> 147
    ws.append(["RECR11.SA", 273, "B3", "Fund", 89.68])    # unchanged (105 + 168)
    ws.append(["SAP.DE", 10, "Frankfurt", "Stock", 98.85])  # non-B3: portfolio-only
    wb.save(path)


@pytest.fixture()
def export_path(tmp_path: Path) -> Path:
    path = tmp_path / "posicao.xlsx"
    _make_b3_export(path)
    return path


@pytest.fixture()
def portfolio_path(tmp_path: Path) -> Path:
    path = tmp_path / "portfolio.xlsx"
    _make_portfolio(path)
    return path


def _rows(path: Path) -> dict[str, float]:
    """Read the portfolio's first sheet back as {ticker: quantity}."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    out = {
        str(r[0]): r[1]
        for r in ws.iter_rows(min_row=2, values_only=True)
        if r[0] is not None
    }
    wb.close()
    return out


def test_parse_aggregates_across_sheets(export_path: Path) -> None:
    entries = _parse_b3_export(export_path)
    # Quantities sum across sheets; Renda Fixa and Total/blank rows are ignored.
    assert entries["ITSA4"].quantity == 9105
    assert entries["FLRY3"].quantity == 2060      # 300 (Acoes) + 1760 (Empréstimos)
    assert entries["BBAS3"].quantity == 1590      # 200 + 1390
    assert entries["KLBN11"].quantity == 2715     # 300 + 2415
    assert entries["RECR11"].quantity == 273      # 105 (Fundo) + 168 (Empréstimos)
    assert "CDB" not in entries                    # Renda Fixa skipped


def test_parse_category_priority(export_path: Path) -> None:
    # RECR11 appears in both Empréstimos (Stock) and Fundo (Fund); Fund wins.
    assert _parse_b3_export(export_path)["RECR11"].category == "Fund"
    assert _parse_b3_export(export_path)["B5P211"].category == "ETF"


async def test_build_plan(export_path: Path, portfolio_path: Path) -> None:
    plan = await build_plan(export_path, portfolio_path)

    updates = {c.ticker: (c.old_quantity, c.new_quantity) for c in plan.updates}
    assert updates == {
        "FLRY3.SA": (1760, 2060),
        "KLBN11.SA": (1515, 2715),
        "BBAS3.SA": (1000, 1590),
        "B5P211.SA": (100, 147),
    }

    # RECR12 is in the export but not the portfolio → an addition, tagged B3/Fund.
    assert len(plan.additions) == 1
    added = plan.additions[0]
    assert added.ticker == "RECR12.SA"
    assert added.quantity == 47
    assert added.exchange == "B3"
    assert added.category == "Fund"

    assert set(plan.unchanged) == {"ITSA4.SA", "RECR11.SA"}
    assert plan.portfolio_only == ["SAP.DE"]   # non-B3 asset never touched
    assert plan.has_changes is True


async def test_apply_plan_updates_only_without_add_new(
    export_path: Path, portfolio_path: Path
) -> None:
    plan = await build_plan(export_path, portfolio_path)
    await apply_plan(plan, portfolio_path, add_new=False)

    rows = _rows(portfolio_path)
    assert rows["FLRY3.SA"] == 2060
    assert rows["KLBN11.SA"] == 2715
    assert rows["BBAS3.SA"] == 1590
    assert rows["B5P211.SA"] == 147
    assert rows["ITSA4.SA"] == 9105       # unchanged
    assert rows["SAP.DE"] == 10           # untouched
    assert "RECR12.SA" not in rows        # not added without --add-new


async def test_apply_plan_with_add_new(
    export_path: Path, portfolio_path: Path
) -> None:
    plan = await build_plan(export_path, portfolio_path)
    await apply_plan(plan, portfolio_path, add_new=True)

    rows = _rows(portfolio_path)
    assert rows["RECR12.SA"] == 47        # now added


async def test_idempotent_second_run(export_path: Path, portfolio_path: Path) -> None:
    # Applying, then re-planning, should show no further quantity updates.
    plan = await build_plan(export_path, portfolio_path)
    await apply_plan(plan, portfolio_path, add_new=False)

    plan2 = await build_plan(export_path, portfolio_path)
    assert plan2.updates == []
