# This module is the counterpart to `reader.py`: where the reader turns the
# Excel spreadsheet into Python objects, the writer turns edits made in the app
# back into the spreadsheet. It is the "write" half of the data pipeline:
#
#   [TUI edit] → writer.py → [Excel file] → reader.py → engine.py → [TUI]
#
# Design goals
# ------------
# 1. **Surgical edits.** The spreadsheet holds data the app does not manage —
#    the hand-typed "Avg Price" column, non-B3 rows, and the whole "FixedIncome"
#    sheet. So we open the existing workbook and change only the specific cells
#    we need, rather than regenerating the file from our in-memory model (which
#    would wipe everything the app doesn't know about).
#
# 2. **Same threading model as the reader.** openpyxl is a blocking library, so
#    the public functions are `async` and run the blocking work in a background
#    thread via `asyncio.to_thread`. This keeps the Textual UI responsive while
#    the file is being written. See `reader.py` for the longer explanation.
#
# 3. **Matches the reader's contract.** The reader looks at the *first* sheet and
#    reads columns A–E as Ticker, Quantity, Exchange, Category, Avg Price. The
#    writer uses those exact positions so a value written here is read back the
#    same way.

import asyncio
from pathlib import Path

import openpyxl
import structlog

from portfolio.core.exceptions import DuplicateTickerError, PositionNotFoundError
from portfolio.core.models import Position

log = structlog.get_logger()

# Column numbers in the positions sheet (openpyxl is 1-indexed: A=1, B=2, ...).
# Keeping them named here means the layout is defined in one place; if the sheet
# ever gains a column, you update these constants instead of hunting for magic
# numbers scattered through the code.
_COL_TICKER = 1     # A
_COL_QUANTITY = 2   # B
_COL_EXCHANGE = 3   # C
_COL_CATEGORY = 4   # D
_COL_AVG_PRICE = 5  # E

_FIRST_DATA_ROW = 2  # row 1 is the header, so real positions start at row 2


def _find_ticker_row(ws: openpyxl.worksheet.worksheet.Worksheet, ticker: str) -> int | None:
    """Return the 1-based row number whose column A matches `ticker`, or None.

    The comparison mirrors the reader: it compares the *stripped* string form of
    each cell so trailing spaces in the spreadsheet never cause a false miss.

    `ws.iter_rows(..., values_only=False)` yields `Cell` objects (not raw values)
    because we need each cell's `.row` number to know where the match is.
    """
    target = ticker.strip()
    for row in ws.iter_rows(min_row=_FIRST_DATA_ROW, min_col=_COL_TICKER, max_col=_COL_TICKER):
        cell = row[0]
        if cell.value is not None and str(cell.value).strip() == target:
            # `int(...)`: openpyxl is untyped, so `cell.row` is seen as `Any`.
            # Converting makes the return type explicitly `int` for the checker.
            return int(cell.row)
    return None


def _last_data_row(ws: openpyxl.worksheet.worksheet.Worksheet) -> int:
    """Return the last row that actually holds a ticker in column A.

    We can't rely on `ws.max_row` because openpyxl counts trailing blank rows
    that Excel sometimes leaves behind. Instead we scan column A and remember the
    highest row that has a value, so a freshly added position lands immediately
    after the real data rather than after a gap of empty rows.
    """
    last = _FIRST_DATA_ROW - 1  # if the sheet has no data rows yet, append at row 2
    for row in ws.iter_rows(min_row=_FIRST_DATA_ROW, min_col=_COL_TICKER, max_col=_COL_TICKER):
        cell = row[0]
        if cell.value is not None and str(cell.value).strip() != "":
            last = int(cell.row)
    return last


def _open_positions_sheet(path: Path) -> tuple[openpyxl.Workbook, openpyxl.worksheet.worksheet.Worksheet]:
    """Open the workbook for editing and return (workbook, first sheet).

    Note the differences from the reader's `load_workbook` call:
      - No `read_only=True`: we need a writable workbook to change cells.
      - No `data_only=True`: `data_only` returns cached values and *drops* any
        formulas on save. Our sheet has no formulas today, but leaving it off is
        the safe default for a writer so we never silently erase them.
    """
    wb = openpyxl.load_workbook(path)
    if not wb.worksheets:
        raise PositionNotFoundError(f"Workbook {path} has no sheets to edit")
    return wb, wb.worksheets[0]


# --- synchronous helpers (blocking; run inside asyncio.to_thread) --------------


def _set_quantity(path: Path, ticker: str, quantity: float) -> None:
    """Update the Quantity cell for an existing ticker. Blocking."""
    if quantity <= 0:
        # Reuse the same rule the `Position` model enforces (`Field(gt=0)`), but
        # check it up front so we fail before touching the file.
        raise ValueError(f"Quantity must be greater than 0, got {quantity}")

    wb, ws = _open_positions_sheet(path)
    row = _find_ticker_row(ws, ticker)
    if row is None:
        wb.close()
        raise PositionNotFoundError(f"Ticker {ticker!r} not found in {path}")

    ws.cell(row=row, column=_COL_QUANTITY, value=quantity)
    wb.save(path)
    wb.close()
    log.info("quantity_updated", ticker=ticker, quantity=quantity)


def _set_avg_price(path: Path, ticker: str, avg_price: float) -> None:
    """Update the Avg Price cell (column E) for an existing ticker. Blocking."""
    if avg_price <= 0:
        # Average price is the denominator of the P&L calculation, so a zero or
        # negative value is meaningless — reject it before touching the file.
        raise ValueError(f"Average price must be greater than 0, got {avg_price}")

    wb, ws = _open_positions_sheet(path)
    row = _find_ticker_row(ws, ticker)
    if row is None:
        wb.close()
        raise PositionNotFoundError(f"Ticker {ticker!r} not found in {path}")

    ws.cell(row=row, column=_COL_AVG_PRICE, value=avg_price)
    wb.save(path)
    wb.close()
    log.info("avg_price_updated", ticker=ticker, avg_price=avg_price)


def _add_position(path: Path, position: Position) -> None:
    """Append a new position row. Blocking. Raises if the ticker already exists."""
    wb, ws = _open_positions_sheet(path)

    if _find_ticker_row(ws, position.ticker) is not None:
        wb.close()
        raise DuplicateTickerError(
            f"Ticker {position.ticker!r} already exists in {path}"
        )

    row = _last_data_row(ws) + 1
    ws.cell(row=row, column=_COL_TICKER, value=position.ticker)
    ws.cell(row=row, column=_COL_QUANTITY, value=position.quantity)
    ws.cell(row=row, column=_COL_EXCHANGE, value=position.exchange)
    ws.cell(row=row, column=_COL_CATEGORY, value=position.category)
    # Only write Avg Price if the user supplied one; leaving the cell blank keeps
    # it consistent with how the reader treats a missing average (P&L = "----").
    if position.avg_price_native is not None:
        ws.cell(row=row, column=_COL_AVG_PRICE, value=position.avg_price_native)

    wb.save(path)
    wb.close()
    log.info("position_added", ticker=position.ticker, quantity=position.quantity)


def _remove_position(path: Path, ticker: str) -> None:
    """Delete the row for an existing ticker. Blocking."""
    wb, ws = _open_positions_sheet(path)
    row = _find_ticker_row(ws, ticker)
    if row is None:
        wb.close()
        raise PositionNotFoundError(f"Ticker {ticker!r} not found in {path}")

    # `delete_rows(idx, amount)` removes `amount` rows starting at `idx` and
    # shifts everything below up by one, so no empty gap is left behind.
    ws.delete_rows(row, 1)
    wb.save(path)
    wb.close()
    log.info("position_removed", ticker=ticker)


# --- async public API (call these from the app) -------------------------------


async def set_quantity(path: Path, ticker: str, quantity: float) -> None:
    """Update the quantity of an existing position, off the event loop."""
    await asyncio.to_thread(_set_quantity, path, ticker, quantity)


async def set_avg_price(path: Path, ticker: str, avg_price: float) -> None:
    """Update the average price of an existing position, off the event loop."""
    await asyncio.to_thread(_set_avg_price, path, ticker, avg_price)


async def add_position(path: Path, position: Position) -> None:
    """Append a new position, off the event loop."""
    await asyncio.to_thread(_add_position, path, position)


async def remove_position(path: Path, ticker: str) -> None:
    """Remove an existing position, off the event loop."""
    await asyncio.to_thread(_remove_position, path, ticker)
